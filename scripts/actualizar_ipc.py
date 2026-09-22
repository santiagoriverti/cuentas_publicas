"""
Actualiza data/reference/IPC.xlsx con los meses nuevos del IPC INDEC.

Uso:
    python scripts/actualizar_ipc.py                      # solo Nivel General (API)
    python scripts/actualizar_ipc.py --divisiones RUTA    # + completa divisiones desde
                                                          #   un IPC.xlsx del mismo formato

- Nivel General: API de Series de Tiempo de datos.gob.ar, serie 148.3_INIVELNAL_DICI_M_26
  (IPC nacional, dic-2016=100). Es la unica columna que usan los notebooks.
- Antes de agregar, compara la API contra los meses que ya estan en el archivo y
  aborta si difieren mas de 0,5 puntos de indice (cambio de base o de serie).
- Divisiones (opcional): se completan solo las celdas vacias, desde el archivo indicado.
"""

import argparse
import io
from pathlib import Path

import openpyxl
import pandas as pd
import requests

ROOT = Path(__file__).resolve().parents[1]
IPC_FILE = ROOT / "data" / "reference" / "IPC.xlsx"
API = "https://apis.datos.gob.ar/series/api/series/"
SERIE = "148.3_INIVELNAL_DICI_M_26"
TOLERANCIA = 0.5


def serie_api() -> pd.Series:
    r = requests.get(API, params={"ids": SERIE, "format": "csv",
                                  "start_date": "2017-01-01", "limit": 5000}, timeout=60)
    r.raise_for_status()
    df = pd.read_csv(io.StringIO(r.text), parse_dates=["indice_tiempo"])
    return df.set_index("indice_tiempo").iloc[:, 0]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--divisiones", help="IPC.xlsx con divisiones para completar celdas vacias")
    args = ap.parse_args()

    api = serie_api()
    wb = openpyxl.load_workbook(IPC_FILE)
    ws = wb.active
    hdr = [c.value for c in ws[1]]

    existentes = {pd.Timestamp(r[0].value): r[1].value
                  for r in ws.iter_rows(min_row=2) if r[0].value is not None}
    comunes = [d for d in existentes if d in api.index]
    dif = max(abs(existentes[d] - api[d]) for d in comunes)
    print(f"Meses comunes archivo/API: {len(comunes)} | dif. maxima: {dif:.4f}")
    if dif > TOLERANCIA:
        raise SystemExit("La API no coincide con el archivo (posible cambio de base). Revisar a mano.")

    ultimo = max(existentes)
    fmt = ws.cell(ws.max_row, 1).number_format
    nuevos = api[api.index > ultimo]
    for d, v in nuevos.items():
        ws.append([d.to_pydatetime(), float(v)])
        ws.cell(ws.max_row, 1).number_format = fmt
        print(f"  agregado {d:%Y-%m}: {v}")
    if nuevos.empty:
        print(f"Sin meses nuevos en la API (ultimo: {ultimo:%Y-%m}).")

    n = 0
    if args.divisiones:
        src = pd.read_excel(args.divisiones).set_index("date")
        for row in ws.iter_rows(min_row=2):
            d = pd.Timestamp(row[0].value)
            for j, c in enumerate(row[1:], 1):
                if c.value is None and d in src.index and hdr[j] in src.columns:
                    c.value = float(src.loc[d, hdr[j]])
                    n += 1
        print(f"Divisiones completadas: {n} celdas")

    cambios = len(nuevos) + (n if args.divisiones else 0)
    if cambios:
        wb.save(IPC_FILE)
        print(f"Guardado: {IPC_FILE.relative_to(ROOT)}")
    else:
        print("Sin cambios: el archivo no se modifico.")


if __name__ == "__main__":
    main()
